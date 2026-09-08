import hashlib
import json
import logging
import os
from dataclasses import asdict, dataclass
from pathlib import Path


BASE_DIR = Path(__file__).resolve().parent.parent
CACHE_DIR = BASE_DIR / "models"
INDEX_DIR = BASE_DIR / "indexes"
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L12-v2"

os.environ.setdefault("HF_HOME", str(CACHE_DIR))
logger = logging.getLogger(__name__)

from langchain_community.document_loaders import PyPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_community.vectorstores.utils import DistanceStrategy
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from transformers import AutoTokenizer


def resolve_model_source(model_name: str = EMBEDDING_MODEL) -> str:
    """Use the bundled snapshot directly and avoid unnecessary network checks."""
    if model_name != EMBEDDING_MODEL:
        return model_name

    model_cache = CACHE_DIR / "models--sentence-transformers--all-MiniLM-L12-v2"
    ref_path = model_cache / "refs" / "main"
    if ref_path.exists():
        revision = ref_path.read_text(encoding="utf-8").strip()
        snapshot = model_cache / "snapshots" / revision
        required_files = (snapshot / "config.json", snapshot / "model.safetensors")
        if all(path.exists() for path in required_files):
            return str(snapshot)
    return model_name


@dataclass(frozen=True)
class SearchResult:
    text: str
    source: str
    page: int
    score: float

    def to_dict(self) -> dict:
        return asdict(self)


class Encoder:
    def __init__(self, model_name: str = EMBEDDING_MODEL, device: str = "cpu"):
        self.embedding_function = HuggingFaceEmbeddings(
            model_name=resolve_model_source(model_name),
            cache_folder=str(CACHE_DIR),
            model_kwargs={"device": device},
            encode_kwargs={"normalize_embeddings": True},
        )


class FaissDb:
    def __init__(self, db: FAISS):
        self.db = db

    @classmethod
    def from_documents(cls, docs: list, embedding_function) -> "FaissDb":
        db = FAISS.from_documents(
            docs,
            embedding_function,
            distance_strategy=DistanceStrategy.COSINE,
        )
        return cls(db)

    @classmethod
    def load(cls, directory: Path, embedding_function) -> "FaissDb":
        # Indexes are generated locally by this app inside a hash-named directory.
        db = FAISS.load_local(
            str(directory),
            embedding_function,
            allow_dangerous_deserialization=True,
            distance_strategy=DistanceStrategy.COSINE,
        )
        return cls(db)

    @property
    def chunk_count(self) -> int:
        return len(self.db.index_to_docstore_id)

    def save(self, directory: Path) -> None:
        directory.mkdir(parents=True, exist_ok=True)
        self.db.save_local(str(directory))

    def search(
        self,
        question: str,
        k: int = 4,
        score_threshold: float = 0.0,
    ) -> list[SearchResult]:
        if not question.strip():
            return []

        # Over-fetch so duplicate/weak passages can be removed without reducing
        # the number of useful sources shown to the user.
        candidates = self.db.similarity_search_with_score(
            question,
            k=max(k * 3, k),
        )
        results = []
        seen_passages = set()
        for doc, distance in candidates:
            # FAISS returns squared L2 distance. Embeddings are normalized, so
            # cosine similarity is 1 - (squared_distance / 2).
            score = max(0.0, min(1.0, 1.0 - (float(distance) / 2.0)))
            if score < score_threshold:
                continue

            passage_key = " ".join(doc.page_content.lower().split())
            if passage_key in seen_passages:
                continue
            seen_passages.add(passage_key)

            source = Path(doc.metadata.get("source", "Uploaded document")).name
            page = int(doc.metadata.get("page", 0)) + 1
            results.append(
                SearchResult(
                    text=doc.page_content.strip(),
                    source=source,
                    page=page,
                    score=score,
                )
            )
            if len(results) == k:
                break

        # If strict filtering dropped all candidates, fall back to the highest scoring available passage(s)
        if not results and candidates:
            for doc, distance in candidates:
                score = max(0.0, min(1.0, 1.0 - (float(distance) / 2.0)))
                passage_key = " ".join(doc.page_content.lower().split())
                if passage_key in seen_passages:
                    continue
                seen_passages.add(passage_key)
                source = Path(doc.metadata.get("source", "Uploaded document")).name
                page = int(doc.metadata.get("page", 0)) + 1
                results.append(
                    SearchResult(
                        text=doc.page_content.strip(),
                        source=source,
                        page=page,
                        score=score,
                    )
                )
                if len(results) == k:
                    break

        return results


def load_single_document(file_path: str) -> list[Document]:
    path = Path(file_path)
    suffix = path.suffix.lower()
    documents = []

    if suffix == ".pdf":
        try:
            return PyPDFLoader(file_path).load()
        except Exception as e:
            logger.warning(f"PyPDFLoader failed for {file_path}: {e}")

    elif suffix in (".docx", ".doc"):
        try:
            import docx

            doc = docx.Document(file_path)
            content_parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    content_parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        content_parts.append(row_text)
            text = "\n\n".join(content_parts)
            if text.strip():
                documents.append(
                    Document(
                        page_content=text.strip(),
                        metadata={"source": file_path, "page": 1},
                    )
                )
            return documents
        except Exception as e:
            logger.warning(f"DOCX loader failed for {file_path}: {e}")

    elif suffix in (".pptx", ".ppt"):
        try:
            import pptx

            prs = pptx.Presentation(file_path)
            for idx, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_texts.append(paragraph.text.strip())
                if slide_texts:
                    documents.append(
                        Document(
                            page_content="\n".join(slide_texts),
                            metadata={"source": file_path, "page": idx},
                        )
                    )
            return documents
        except Exception as e:
            logger.warning(f"PPTX loader failed for {file_path}: {e}")

    elif suffix in (".xlsx", ".xls"):
        try:
            import pandas as pd

            with pd.ExcelFile(file_path) as excel_file:
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet_name)
                    if not df.empty:
                        sheet_text = (
                            f"Sheet: {sheet_name}\n" + df.to_string(index=False)
                        )
                        documents.append(
                            Document(
                                page_content=sheet_text,
                                metadata={"source": file_path, "page": sheet_name},
                            )
                        )
            return documents
        except Exception as e:
            logger.warning(f"Excel loader failed for {file_path}: {e}")

    elif suffix in (".csv", ".tsv"):
        try:
            import pandas as pd

            sep = "\t" if suffix == ".tsv" else ","
            df = pd.read_csv(file_path, sep=sep)
            if not df.empty:
                text = df.to_string(index=False)
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_path, "page": 1},
                    )
                )
                return documents
        except Exception as e:
            logger.warning(f"CSV loader failed for {file_path}: {e}")

    # For text-based formats (.txt, .md, .json, .xml, .html, .py, .js, .sql, .yaml, .log, etc.) or generic text fallback
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            text = path.read_text(encoding=encoding, errors="replace")
            # Make sure it's not a pure binary file with null bytes
            if text.strip() and "\x00" not in text[:500]:
                documents.append(
                    Document(
                        page_content=text.strip(),
                        metadata={"source": file_path, "page": 1},
                    )
                )
                return documents
        except Exception:
            continue

    return documents


def load_and_split_documents(
    file_paths: list[str],
    chunk_size: int = 600,
    chunk_overlap: int = 100,
) -> list:
    pages = []
    for file_path in file_paths:
        pages.extend(load_single_document(file_path))

    tokenizer = AutoTokenizer.from_pretrained(
        resolve_model_source(),
        cache_dir=str(CACHE_DIR),
    )
    text_splitter = RecursiveCharacterTextSplitter.from_huggingface_tokenizer(
        tokenizer=tokenizer,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        strip_whitespace=True,
    )
    return text_splitter.split_documents(pages)


# Alias for backward compatibility
load_and_split_pdfs = load_and_split_documents


def collection_fingerprint(
    files: list[tuple[str, bytes]],
    chunk_size: int,
    chunk_overlap: int,
) -> str:
    digest = hashlib.sha256()
    digest.update(f"{EMBEDDING_MODEL}:{chunk_size}:{chunk_overlap}".encode())
    for name, content in sorted(files, key=lambda item: item[0].lower()):
        digest.update(name.encode("utf-8", errors="replace"))
        digest.update(len(content).to_bytes(8, "big"))
        digest.update(content)
    return digest.hexdigest()[:20]


def index_path(collection_id: str) -> Path:
    return INDEX_DIR / collection_id


def save_manifest(directory: Path, manifest: dict) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    with (directory / "manifest.json").open("w", encoding="utf-8") as file:
        json.dump(manifest, file, indent=2)


def load_manifest(directory: Path) -> dict:
    manifest_path = directory / "manifest.json"
    if not manifest_path.exists():
        return {}
    with manifest_path.open(encoding="utf-8") as file:
        return json.load(file)
